import json
import xml.etree.ElementTree as ET
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from experttwin import providers
from experttwin.config import Settings
from experttwin.db import Database
from experttwin.parsers import parse_file
from experttwin.providers import (
    AzureOpenAIAnswerer,
    AzureOpenAIDecisionExtractor,
    HeuristicDecisionExtractor,
    LocalUnavailableTranscriber,
)
from experttwin.services import ChatService, FingerprintService, IngestionService, RetrievalService

WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_2010_NAMESPACE = "http://schemas.microsoft.com/office/word/2010/wordml"
WORD_2012_NAMESPACE = "http://schemas.microsoft.com/office/word/2012/wordml"
RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPE_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/content-types"
W = f"{{{WORD_NAMESPACE}}}"
W14 = f"{{{WORD_2010_NAMESPACE}}}"
W15 = f"{{{WORD_2012_NAMESPACE}}}"


def _comment(
    comment_id: str,
    author: str,
    text: str,
    para_id: str,
    date: str | None = None,
) -> ET.Element:
    attributes = {f"{W}id": comment_id, f"{W}author": author}
    if date:
        attributes[f"{W}date"] = date
    comment = ET.Element(f"{W}comment", attributes)
    paragraph = ET.SubElement(comment, f"{W}p", {f"{W14}paraId": para_id})
    run = ET.SubElement(paragraph, f"{W}r")
    ET.SubElement(run, f"{W}t").text = text
    return comment


def _reference_run(comment_id: str) -> ET.Element:
    run = ET.Element(f"{W}r")
    ET.SubElement(run, f"{W}commentReference", {f"{W}id": comment_id})
    return run


def _add_comment_relationship(parts: dict[str, bytes], include_threading: bool) -> None:
    relationships_path = "word/_rels/document.xml.rels"
    relationships = ET.fromstring(parts[relationships_path])
    existing_ids = {item.get("Id") for item in relationships}
    relationship_id = "rIdComments"
    while relationship_id in existing_ids:
        relationship_id += "X"
    ET.SubElement(
        relationships,
        f"{{{RELATIONSHIP_NAMESPACE}}}Relationship",
        {
            "Id": relationship_id,
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/comments"
            ),
            "Target": "comments.xml",
        },
    )
    if include_threading:
        ET.SubElement(
            relationships,
            f"{{{RELATIONSHIP_NAMESPACE}}}Relationship",
            {
                "Id": f"{relationship_id}Extended",
                "Type": (
                    "http://schemas.microsoft.com/office/2011/"
                    "relationships/commentsExtended"
                ),
                "Target": "commentsExtended.xml",
            },
        )
    parts[relationships_path] = ET.tostring(
        relationships, encoding="utf-8", xml_declaration=True
    )

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(
        content_types,
        f"{{{CONTENT_TYPE_NAMESPACE}}}Override",
        {
            "PartName": "/word/comments.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.comments+xml"
            ),
        },
    )
    if include_threading:
        ET.SubElement(
            content_types,
            f"{{{CONTENT_TYPE_NAMESPACE}}}Override",
            {
                "PartName": "/word/commentsExtended.xml",
                "ContentType": "application/vnd.ms-word.commentsExtended+xml",
            },
        )
    parts["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )


def create_synthetic_docx(
    path: Path,
    include_comments: bool = True,
    include_threading: bool = True,
) -> None:
    document = Document()
    document.add_paragraph("A neutral body paragraph remains available.")
    document.add_paragraph("A second body paragraph provides nearby context.")
    document.save(path)
    if not include_comments:
        return

    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    document_xml = ET.fromstring(parts["word/document.xml"])
    paragraphs = list(document_xml.iter(f"{W}p"))
    first_run = next(paragraphs[0].iter(f"{W}r"))
    first_run_index = list(paragraphs[0]).index(first_run)
    paragraphs[0].insert(first_run_index, ET.Element(f"{W}commentRangeStart", {f"{W}id": "8"}))
    paragraphs[0].insert(
        first_run_index + 2, ET.Element(f"{W}commentRangeEnd", {f"{W}id": "8"})
    )
    paragraphs[0].insert(first_run_index + 3, _reference_run("8"))
    paragraphs[1].append(_reference_run("9"))
    parts["word/document.xml"] = ET.tostring(
        document_xml, encoding="utf-8", xml_declaration=True
    )

    comments = ET.Element(f"{W}comments")
    comments.append(
        _comment(
            "7",
            "aShA",
            "We should choose queued writes because they isolate retries.",
            "A0000007",
            "2026-01-02T03:04:05Z",
        )
    )
    comments.append(
        _comment(
            "8",
            "Other Reviewer",
            "We chose direct writes because they look simpler.",
            "B0000008",
        )
    )
    comments.append(
        _comment(
            "9",
            "ASHA",
            "Prefer bounded retries over unlimited retries because failures must surface.",
            "A0000009",
        )
    )
    parts["word/comments.xml"] = ET.tostring(
        comments, encoding="utf-8", xml_declaration=True
    )
    if include_threading:
        extended = ET.Element(f"{W15}commentsEx")
        ET.SubElement(
            extended,
            f"{W15}commentEx",
            {
                f"{W15}paraId": "A0000007",
                f"{W15}paraIdParent": "B0000008",
            },
        )
        ET.SubElement(
            extended,
            f"{W15}commentEx",
            {f"{W15}paraId": "A0000009"},
        )
        parts["word/commentsExtended.xml"] = ET.tostring(
            extended, encoding="utf-8", xml_declaration=True
        )
    _add_comment_relationship(parts, include_threading)

    rebuilt = path.with_suffix(".rebuilt.docx")
    with ZipFile(rebuilt, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    rebuilt.replace(path)


def test_docx_body_and_matching_comments_are_preserved(runtime_dir: Path):
    path = runtime_dir / "review.docx"
    create_synthetic_docx(path)

    segments = parse_file(path, reviewer_name="Asha")
    body = [segment for segment in segments if segment.comment_id is None]
    comments = [segment for segment in segments if segment.comment_id is not None]

    assert body
    assert "neutral body paragraph" in body[0].text
    assert {comment.comment_id for comment in comments} == {"7", "9"}
    assert comments[0].author == "aShA"
    assert comments[0].timestamp == "2026-01-02T03:04:05Z"
    assert comments[0].anchor_text == "A neutral body paragraph remains available."
    assert comments[0].parent_comment_id == "8"
    assert comments[0].parent_author == "Other Reviewer"
    assert comments[0].parent_text == "We chose direct writes because they look simpler."
    assert "Word comment 7" in comments[0].page_or_segment
    assert "reply to comment 8 by Other Reviewer" in comments[0].page_or_segment
    assert "paragraph 1" in comments[0].page_or_segment
    assert comments[1].anchor_text == "A second body paragraph provides nearby context."
    assert all("direct writes" not in comment.text for comment in comments)


def test_docx_without_comments_keeps_body_text(runtime_dir: Path):
    path = runtime_dir / "plain.docx"
    create_synthetic_docx(path, include_comments=False)

    segments = parse_file(path, reviewer_name="Asha")

    assert len(segments) == 1
    assert segments[0].comment_id is None
    assert "neutral body paragraph" in segments[0].text


def test_docx_without_comments_extended_keeps_comments(runtime_dir: Path):
    path = runtime_dir / "unthreaded-review.docx"
    create_synthetic_docx(path, include_threading=False)

    comments = [
        segment
        for segment in parse_file(path, reviewer_name="Asha")
        if segment.comment_id is not None
    ]

    assert {comment.comment_id for comment in comments} == {"7", "9"}
    assert all(comment.parent_comment_id is None for comment in comments)


def test_ingestion_filters_comments_by_selected_expert(runtime_dir: Path):
    path = runtime_dir / "review.docx"
    create_synthetic_docx(path)
    database = Database(runtime_dir / "comments.db")
    database.initialize()
    expert = database.create_expert("asha")

    result = IngestionService(
        database,
        HeuristicDecisionExtractor(),
        LocalUnavailableTranscriber(),
    ).ingest(expert.id, path, "Synthetic review")

    assert len(result.decisions) == 2
    assert all("direct writes" not in decision.choice for decision in result.decisions)
    assert all(decision.evidence_type == "review_comment" for decision in result.decisions)
    assert all(decision.evidence_author for decision in result.decisions)
    assert all(
        "Anchored document context:" in decision.evidence_text
        for decision in result.decisions
    )
    threaded = next(decision for decision in result.decisions if decision.comment_id == "7")
    assert threaded.parent_comment_id == "8"
    assert threaded.parent_author == "Other Reviewer"
    assert "Parent review comment by Other Reviewer:" in threaded.evidence_text
    assert "Selected expert reply by aShA:" in threaded.evidence_text
    assert {decision.page_or_segment.split(" — ")[0] for decision in result.decisions} == {
        "Word comment 7",
        "Word comment 9",
    }
    fingerprint = FingerprintService(database).derive(expert.id)
    assert fingerprint.expert_attributed_decision_count == 2
    assert fingerprint.contextual_decision_count == 0


class _FakeCompletions:
    def __init__(self):
        self.prompts: list[str] = []
        self.requests: list[dict] = []

    def create(self, **kwargs):
        self.requests.append(kwargs)
        prompt = kwargs["messages"][-1]["content"]
        self.prompts.append(prompt)
        decisions = []
        if "Evidence type:" in prompt:
            decisions.append(
                {
                    "decision": "Queue writes for retry isolation.",
                    "choice": "Use queued writes",
                    "alternatives": ["Direct writes"],
                    "rationale": ["Retries are isolated"],
                    "constraints": [],
                    "risks": ["Queue backlog"],
                    "outcome": None,
                    "tags": ["reliability"],
                    "confidence": 0.84,
                }
            )
        message = SimpleNamespace(content=json.dumps({"decisions": decisions}))
        return SimpleNamespace(choices=[SimpleNamespace(message=message)])


def test_azure_comment_prompt_and_attribution_survive_mapping(
    runtime_dir: Path, monkeypatch
):
    path = runtime_dir / "review.docx"
    create_synthetic_docx(path)
    comment = next(
        segment
        for segment in parse_file(path, reviewer_name="Asha")
        if segment.comment_id == "7"
    )
    completions = _FakeCompletions()
    fake_client = SimpleNamespace(
        chat=SimpleNamespace(completions=completions)
    )
    monkeypatch.setattr(providers, "_azure_openai_client", lambda _settings: fake_client)
    extractor = AzureOpenAIDecisionExtractor(
        Settings(
            azure_openai_endpoint="https://example.invalid/",
            azure_openai_chat_deployment="test-deployment",
        )
    )

    database = Database(runtime_dir / "azure-comments.db")
    database.initialize()
    expert = database.create_expert("Asha")
    source = database.create_source(expert.id, "Synthetic review", path.name, "document")
    decisions = extractor.extract(
        expert.id,
        expert.name,
        source.id,
        source.title,
        [comment],
    )
    database.store_decisions(decisions)
    mapped = database.list_decisions(expert.id)[0]

    prompt = completions.prompts[0]
    assert "temperature" not in completions.requests[0]
    assert completions.requests[0]["response_format"] == {"type": "json_object"}
    assert "Evidence type: threaded review reply" in prompt
    assert "Reviewer: aShA" in prompt
    assert "Selected expert: Asha" in prompt
    assert "Comment ID: 7" in prompt
    assert "Parent comment ID: 8" in prompt
    assert "Parent reviewer: Other Reviewer" in prompt
    assert "context only; do not attribute" in prompt
    assert comment.text in prompt
    assert comment.parent_text in prompt
    assert comment.anchor_text in prompt
    assert "recommendations, objections, approvals, risk concerns" in prompt
    assert "acceptance, rejection, refinement, answers, risk-gap acknowledgment" in prompt
    assert mapped.evidence_type == "review_comment"
    assert mapped.evidence_author == "aShA"
    assert mapped.comment_id == "7"
    assert mapped.anchor_text == comment.anchor_text
    assert mapped.parent_comment_id == "8"
    assert mapped.parent_author == "Other Reviewer"
    assert mapped.parent_text == comment.parent_text
    assert "Parent review comment by Other Reviewer:" in mapped.evidence_text
    assert "Selected expert reply by aShA:" in mapped.evidence_text
    assert "Anchored document context:" in mapped.evidence_text

    response = ChatService(
        RetrievalService(database),
        FingerprintService(database),
    ).answer(expert.id, "queued writes and retries")
    assert response.citations[0].evidence_type == "review_comment"
    assert response.citations[0].evidence_author == "aShA"
    assert response.citations[0].comment_id == "7"
    assert response.citations[0].anchor_text == comment.anchor_text
    assert response.citations[0].parent_comment_id == "8"
    assert response.citations[0].parent_author == "Other Reviewer"

    ui_script = Path("src/experttwin/static/app.js").read_text(encoding="utf-8")
    assert "decision.evidence_author" in ui_script
    assert "citation.parent_comment_id" in ui_script
    assert "citation.parent_author" in ui_script
    assert "Unattributed context" in ui_script
    assert "expert_attributed_decision_count" in ui_script
    assert "attribution_status" in ui_script


def test_azure_answerer_does_not_force_temperature(monkeypatch):
    requests = []

    def create(**kwargs):
        requests.append(kwargs)
        message = SimpleNamespace(content="Grounded answer")
        return SimpleNamespace(choices=[SimpleNamespace(message=message)])

    fake_client = SimpleNamespace(
        chat=SimpleNamespace(completions=SimpleNamespace(create=create))
    )
    monkeypatch.setattr(providers, "_azure_openai_client", lambda _settings: fake_client)
    answerer = AzureOpenAIAnswerer(
        Settings(
            azure_openai_endpoint="https://example.invalid/",
            azure_openai_chat_deployment="test-deployment",
        )
    )

    answer = answerer.answer("Question?", "[1] Evidence", "Pattern")

    assert answer == "Grounded answer"
    assert requests[0]["model"] == "test-deployment"
    assert "temperature" not in requests[0]
    system_prompt = requests[0]["messages"][0]["content"]
    assert "EXPERT-ATTRIBUTED JUDGMENT" in system_prompt
    assert "CONTEXTUAL DOCUMENT CONTENT" in system_prompt
    assert "Never claim contextual content is the expert's view" in system_prompt
