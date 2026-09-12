import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile

from .models import ParsedSegment

WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_2010_NAMESPACE = "http://schemas.microsoft.com/office/word/2010/wordml"
WORD_2012_NAMESPACE = "http://schemas.microsoft.com/office/word/2012/wordml"
W = f"{{{WORD_NAMESPACE}}}"
W14 = f"{{{WORD_2010_NAMESPACE}}}"
W15 = f"{{{WORD_2012_NAMESPACE}}}"
TRANSCRIPT_TURN_RE = re.compile(
    r"^\[(?P<timestamp>(?:\d{4}-\d{2}-\d{2}\s+\d{1,2}:\d{2}|"
    r"\d{1,2}:\d{2}(?::\d{2})?))\]\s+"
    r"(?P<speaker>[^:\n]+?)(?:\s+\((?P<role>[^)\n]+)\))?:\s*(?P<text>.*)$"
)


class UnsupportedFileError(ValueError):
    pass


def _chunk_text(text: str, prefix: str = "segment", max_chars: int = 2200) -> list[ParsedSegment]:
    text = text.replace("\x00", "").strip()
    if not text:
        return []
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if current and len(current) + len(paragraph) + 2 > max_chars:
            chunks.append(current)
            current = paragraph
        else:
            current = f"{current}\n\n{paragraph}".strip()
    if current:
        chunks.append(current)
    return [
        ParsedSegment(text=chunk, page_or_segment=f"{prefix} {index}")
        for index, chunk in enumerate(chunks, start=1)
    ]


def _parse_named_speaker_transcript(text: str) -> list[ParsedSegment] | None:
    turns = []
    current = None
    preamble = []

    def flush() -> None:
        nonlocal current
        if not current:
            return
        turn_text = "\n".join(current["lines"]).strip()
        if turn_text:
            timestamp = current["timestamp"]
            speaker = current["speaker"]
            role = current["role"]
            turn_kind = "chat" if "-" in timestamp else "meeting"
            role_label = f" ({role})" if role else ""
            turns.append(
                ParsedSegment(
                    text=turn_text,
                    page_or_segment=(
                        f"{turn_kind} turn at {timestamp} — {speaker}{role_label}"
                    ),
                    timestamp=timestamp,
                    author=speaker,
                    speaker_role=role,
                )
            )
        current = None

    for line in text.replace("\x00", "").splitlines():
        match = TRANSCRIPT_TURN_RE.match(line.strip())
        if match:
            flush()
            current = {
                "timestamp": match.group("timestamp"),
                "speaker": match.group("speaker").strip(),
                "role": (match.group("role") or "").strip() or None,
                "lines": [match.group("text").strip()],
            }
        elif current is not None:
            if line.strip():
                current["lines"].append(line.strip())
        elif line.strip():
            preamble.append(line.strip())
    flush()
    if not turns:
        return None
    if preamble:
        turns.insert(
            0,
            ParsedSegment(
                text="\n".join(preamble),
                page_or_segment="transcript context",
            ),
        )
    return turns


def _element_text(element: ET.Element) -> str:
    paragraphs = []
    for paragraph in element.findall(f".//{W}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{W}t")).strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs).strip()


def _docx_comment_anchors(document_xml: bytes) -> dict[str, tuple[str, int]]:
    root = ET.fromstring(document_xml)
    anchors: dict[str, list[str]] = defaultdict(list)
    context: dict[str, tuple[str, int]] = {}
    active: set[str] = set()

    for paragraph_number, paragraph in enumerate(root.iter(f"{W}p"), start=1):
        paragraph_text = "".join(node.text or "" for node in paragraph.iter(f"{W}t")).strip()
        for node in paragraph.iter():
            if node.tag == f"{W}commentRangeStart":
                comment_id = node.get(f"{W}id")
                if comment_id is not None:
                    active.add(comment_id)
                    context.setdefault(comment_id, (paragraph_text, paragraph_number))
            elif node.tag == f"{W}t" and node.text:
                for comment_id in active:
                    if sum(len(part) for part in anchors[comment_id]) < 800:
                        anchors[comment_id].append(node.text)
            elif node.tag == f"{W}commentRangeEnd":
                comment_id = node.get(f"{W}id")
                if comment_id is not None:
                    active.discard(comment_id)
            elif node.tag == f"{W}commentReference":
                comment_id = node.get(f"{W}id")
                if comment_id is not None:
                    context.setdefault(comment_id, (paragraph_text, paragraph_number))

    result = {}
    for comment_id in set(anchors) | set(context):
        anchor = "".join(anchors.get(comment_id, [])).strip()
        nearby, paragraph_number = context.get(comment_id, ("", 0))
        result[comment_id] = ((anchor or nearby)[:500], paragraph_number)
    return result


def _parse_docx_comments(path: Path, reviewer_name: str | None) -> list[ParsedSegment]:
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/comments.xml" not in names:
            return []
        comments_root = ET.fromstring(archive.read("word/comments.xml"))
        parent_para_ids = {}
        if "word/commentsExtended.xml" in names:
            extended_root = ET.fromstring(archive.read("word/commentsExtended.xml"))
            parent_para_ids = {
                para_id.casefold(): parent_id.casefold()
                for comment in extended_root.iter(f"{W15}commentEx")
                if (para_id := comment.get(f"{W15}paraId"))
                and (parent_id := comment.get(f"{W15}paraIdParent"))
            }
        anchors = (
            _docx_comment_anchors(archive.read("word/document.xml"))
            if "word/document.xml" in names
            else {}
        )

    comment_records = []
    comments_by_para_id = {}
    for comment in comments_root.iter(f"{W}comment"):
        paragraph = next(comment.iter(f"{W}p"), None)
        record = {
            "id": comment.get(f"{W}id") or "unknown",
            "author": (comment.get(f"{W}author") or "").strip(),
            "date": comment.get(f"{W}date"),
            "text": _element_text(comment),
            "para_id": paragraph.get(f"{W14}paraId") if paragraph is not None else None,
        }
        comment_records.append(record)
        if record["para_id"]:
            comments_by_para_id[record["para_id"].casefold()] = record

    expected_author = reviewer_name.strip().casefold() if reviewer_name else None
    segments = []
    for record in comment_records:
        author = record["author"]
        if expected_author is not None and author.casefold() != expected_author:
            continue
        text = record["text"]
        if not text:
            continue
        comment_id = record["id"]
        date = record["date"]
        para_id = record["para_id"].casefold() if record["para_id"] else ""
        parent = comments_by_para_id.get(parent_para_ids.get(para_id, ""))
        anchor_text, paragraph_number = anchors.get(comment_id, ("", 0))
        if not anchor_text and parent:
            anchor_text, paragraph_number = anchors.get(parent["id"], ("", 0))
        location_parts = [f"Word comment {comment_id}", f"author {author or 'unknown'}"]
        if parent:
            location_parts.append(
                f"reply to comment {parent['id']} by {parent['author'] or 'unknown'}"
            )
        if paragraph_number:
            location_parts.append(f"paragraph {paragraph_number}")
        if anchor_text:
            location_parts.append(f'anchor "{anchor_text}"')
        else:
            location_parts.append("anchor unavailable")
        segments.append(
            ParsedSegment(
                text=text,
                page_or_segment=" — ".join(location_parts),
                timestamp=date,
                author=author or None,
                comment_id=comment_id,
                anchor_text=anchor_text or None,
                parent_comment_id=parent["id"] if parent else None,
                parent_author=(parent["author"] or None) if parent else None,
                parent_text=(parent["text"] or None) if parent else None,
            )
        )
    return segments


def parse_file(path: Path, reviewer_name: str | None = None) -> list[ParsedSegment]:
    extension = path.suffix.lower()
    if extension == ".txt":
        text = path.read_text(encoding="utf-8-sig")
        return _parse_named_speaker_transcript(text) or _chunk_text(text, "segment")
    if extension == ".md":
        return _chunk_text(path.read_text(encoding="utf-8-sig"), "segment")
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [
            ParsedSegment(text=text, page_or_segment=f"page {number}")
            for number, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        ]
    if extension == ".docx":
        from docx import Document

        document = Document(str(path))
        body_segments = _chunk_text(
            "\n\n".join(p.text for p in document.paragraphs if p.text.strip()), "section"
        )
        return body_segments + _parse_docx_comments(path, reviewer_name)
    if extension == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        segments = []
        for number, slide in enumerate(presentation.slides, start=1):
            text = "\n".join(
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            )
            if text:
                segments.append(ParsedSegment(text=text, page_or_segment=f"slide {number}"))
        return segments
    raise UnsupportedFileError(
        f"Unsupported document type '{extension}'. Upload TXT, MD, PDF, DOCX, or PPTX."
    )
