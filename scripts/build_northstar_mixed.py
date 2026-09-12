"""Build the fictional Project Northstar mixed-artifact corpus."""

import json
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from experttwin.parsers import parse_file
from experttwin.role_simulations import ROSTER

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DATA = ROOT / "demo-data" / "role-simulations"
OUTPUT = SOURCE_DATA / "northstar-mixed"
NOTICE = (
    "SYNTHETIC ROLE-BASED SIMULATION — NOT BASED ON THESE EMPLOYEES' "
    "ACTUAL BEHAVIOR OR WORK HISTORY."
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
W15_NS = "http://schemas.microsoft.com/office/word/2012/wordml"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
W = f"{{{W_NS}}}"
W14 = f"{{{W14_NS}}}"
W15 = f"{{{W15_NS}}}"

ARTIFACTS = [
    ("01-initial-design-review.docx", "Initial design review with Word comments", "document"),
    ("02-group-chat.txt", "Timestamped group design chat", "meeting-transcript"),
    (
        "03-architecture-review-transcript.txt",
        "Architecture review accept-reject-refine transcript",
        "meeting-transcript",
    ),
    ("04-architecture-review.wav", "Architecture review audio rendition", "meeting-recording"),
    ("05-incident-report.md", "Projection lag incident report", "document"),
    ("06-revised-adr.md", "Revised architecture decision record", "document"),
    ("07-measured-outcome.md", "Measured outcome and retrospective", "document"),
]


def _profile_records() -> list[dict]:
    records = []
    for path in sorted(SOURCE_DATA.glob("*.json")):
        record = json.loads(path.read_text(encoding="utf-8"))
        if record.get("evidence_author") in ROSTER:
            records.append(record)
    by_name = {record["evidence_author"]: record for record in records}
    if set(by_name) != set(ROSTER):
        raise ValueError("The base role-simulation records do not match the roster.")
    return [by_name[name] for name in ROSTER]


def _comment(comment_id: int, author: str, text: str, para_id: str) -> ET.Element:
    node = ET.Element(
        f"{W}comment",
        {
            f"{W}id": str(comment_id),
            f"{W}author": author,
            f"{W}date": f"2026-09-12T{8 + comment_id // 2:02d}:00:00Z",
        },
    )
    paragraph = ET.SubElement(node, f"{W}p", {f"{W14}paraId": para_id})
    run = ET.SubElement(paragraph, f"{W}r")
    ET.SubElement(run, f"{W}t").text = text
    return node


def _add_reference(paragraph: ET.Element, comment_id: int) -> None:
    run = ET.Element(f"{W}r")
    ET.SubElement(run, f"{W}commentReference", {f"{W}id": str(comment_id)})
    paragraph.append(run)


def _add_comment_parts(parts: dict[str, bytes]) -> None:
    relationships_path = "word/_rels/document.xml.rels"
    relationships = ET.fromstring(parts[relationships_path])
    ET.SubElement(
        relationships,
        f"{{{REL_NS}}}Relationship",
        {
            "Id": "rIdNorthstarComments",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/comments"
            ),
            "Target": "comments.xml",
        },
    )
    ET.SubElement(
        relationships,
        f"{{{REL_NS}}}Relationship",
        {
            "Id": "rIdNorthstarCommentsExtended",
            "Type": (
                "http://schemas.microsoft.com/office/2011/"
                "relationships/commentsExtended"
            ),
            "Target": "commentsExtended.xml",
        },
    )
    parts[relationships_path] = ET.tostring(
        relationships, encoding="utf-8", xml_declaration=True
    )

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Override",
        {
            "PartName": "/word/comments.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.comments+xml"
            ),
        },
    )
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Override",
        {
            "PartName": "/word/commentsExtended.xml",
            "ContentType": "application/vnd.ms-word.commentsExtended+xml",
        },
    )
    parts["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )


def _build_docx(records: list[dict]) -> Path:
    path = OUTPUT / ARTIFACTS[0][0]
    document = Document()
    document.add_heading(f"{NOTICE} Project Northstar Initial Design Review", level=1)
    document.add_paragraph(
        f"{NOTICE} Project Northstar is a fictional 48-hour exercise for imaginary polar "
        "research stations. All comments, metrics, systems, and outcomes are invented."
    )
    anchors = [
        document.add_paragraph(
            "Scope and release gates: demonstrate a safe synthetic beacon-to-dashboard "
            "workflow under intermittent connectivity."
        ),
        document.add_paragraph(
            "Event contract and command handling: preserve replay safety, schema clarity, "
            "and predictable failure behavior."
        ),
        document.add_paragraph(
            "Service architecture: separate command acceptance from projections and keep "
            "external dependencies replaceable."
        ),
        document.add_paragraph(
            "Data path: validate, quarantine, reconcile, and retain lineage for every "
            "synthetic batch."
        ),
        document.add_paragraph(
            "Operator experience: expose freshness, accessible recovery controls, and "
            "actionable error references."
        ),
        document.add_paragraph(
            "Operational readiness: use bounded concurrency, measurable latency budgets, "
            "and deterministic recovery drills."
        ),
    ]
    document.save(path)

    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    document_xml = ET.fromstring(parts["word/document.xml"])
    xml_paragraphs = list(document_xml.iter(f"{W}p"))
    anchor_nodes = []
    for anchor in anchors:
        anchor_nodes.append(
            next(
                paragraph
                for paragraph in xml_paragraphs
                if "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
                == anchor.text
            )
        )

    comments = ET.Element(f"{W}comments")
    extended = ET.Element(f"{W15}commentsEx")
    names = list(ROSTER)
    thread_pairs = [
        (names[0], names[1]),
        (names[2], names[3]),
        (names[4], names[5]),
        (names[6], names[7]),
        (names[8], names[9]),
        (names[10], names[0]),
    ]
    record_by_name = {record["evidence_author"]: record for record in records}
    comment_id = 1
    for anchor_index, (root_author, reply_author) in enumerate(thread_pairs):
        root_id = comment_id
        reply_id = comment_id + 1
        root_text = (
            f"{NOTICE} {record_by_name[root_author]['decisions'][0]['statement']} "
            f"Choice: {record_by_name[root_author]['decisions'][0]['choice']}."
        )
        reply_text = (
            f"{NOTICE} Replying to the fictional design thread: "
            f"{record_by_name[reply_author]['decisions'][0]['statement']} "
            f"Choice: {record_by_name[reply_author]['decisions'][0]['choice']}."
        )
        paragraph = anchor_nodes[anchor_index]
        first_run = next(paragraph.iter(f"{W}r"))
        index = list(paragraph).index(first_run)
        paragraph.insert(
            index, ET.Element(f"{W}commentRangeStart", {f"{W}id": str(root_id)})
        )
        paragraph.insert(
            index + 2,
            ET.Element(f"{W}commentRangeEnd", {f"{W}id": str(root_id)}),
        )
        _add_reference(paragraph, root_id)

        root_para_id = f"A{root_id:07X}"
        reply_para_id = f"B{reply_id:07X}"
        comments.append(_comment(root_id, root_author, root_text, root_para_id))
        comments.append(_comment(reply_id, reply_author, reply_text, reply_para_id))
        ET.SubElement(extended, f"{W15}commentEx", {f"{W15}paraId": root_para_id})
        ET.SubElement(
            extended,
            f"{W15}commentEx",
            {
                f"{W15}paraId": reply_para_id,
                f"{W15}paraIdParent": root_para_id,
            },
        )
        comment_id += 2

    parts["word/document.xml"] = ET.tostring(
        document_xml, encoding="utf-8", xml_declaration=True
    )
    parts["word/comments.xml"] = ET.tostring(
        comments, encoding="utf-8", xml_declaration=True
    )
    parts["word/commentsExtended.xml"] = ET.tostring(
        extended, encoding="utf-8", xml_declaration=True
    )
    _add_comment_parts(parts)
    rebuilt = path.with_suffix(".rebuilt.docx")
    with ZipFile(rebuilt, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    rebuilt.replace(path)
    return path


def _write_group_chat(records: list[dict]) -> Path:
    path = OUTPUT / ARTIFACTS[1][0]
    lines = [
        NOTICE,
        "Project Northstar and all messages, tradeoffs, metrics, and decisions are fictional.",
        "",
    ]
    for index, record in enumerate(records):
        decision = record["decisions"][1]
        minute = 5 + index * 3
        lines.extend(
            [
                (
                    f"[2026-09-12 09:{minute:02d}] {record['evidence_author']} "
                    f"({record['evidence_role']}): {NOTICE}"
                ),
                decision["statement"],
                (
                    f"I prefer {decision['choice']} rather than "
                    f"{'; '.join(decision['alternatives'])}."
                ),
                (
                    f"Constraint: {'; '.join(decision['constraints'])}. "
                    f"Risk: {'; '.join(decision['risks'])}."
                ),
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _write_architecture_transcript(records: list[dict]) -> Path:
    path = OUTPUT / ARTIFACTS[2][0]
    verbs = ["Accepted", "Refined", "Accepted", "Accepted", "Rejected", "Refined"]
    lines = [
        NOTICE,
        "This Project Northstar architecture review is wholly fictional.",
        "",
    ]
    for index, record in enumerate(records):
        decision = record["decisions"][2]
        seconds = index * 24
        timestamp = f"{seconds // 60:02d}:{seconds % 60:02d}"
        verb = verbs[index % len(verbs)]
        lines.extend(
            [
                (
                    f"[{timestamp}] {record['evidence_author']} "
                    f"({record['evidence_role']}): {NOTICE} {verb}: "
                    f"{decision['choice']}."
                ),
                (
                    f"Rationale: {'; '.join(decision['rationale'])}. "
                    f"Alternative: {'; '.join(decision['alternatives'])}."
                ),
                "",
            ]
        )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _write_markdown(records: list[dict]) -> None:
    incident = f"""# {NOTICE}

# Project Northstar Synthetic Projection-Lag Incident

## Impact
For 17 fictional minutes, dashboard projections were up to 94 seconds behind while
command acceptance remained available. No real systems, people, or data were involved.

## Timeline
- 14:02 — A synthetic burst increased queue depth.
- 14:06 — The freshness indicator crossed its 60-second warning threshold.
- 14:09 — Operators used stage metrics to isolate an intentionally constrained worker.
- 14:19 — Bounded concurrency was raised from four to eight after a replay-safe checkpoint.

## Root cause
The exercise configured worker concurrency below the measured arrival rate. The queue,
idempotency keys, and deterministic checkpoints behaved as designed.

## Decisions
- Preserve command acceptance and drain the durable queue rather than bypass validation.
- Keep the stale-data warning visible and report queue age separately from request latency.
- Add a release gate requiring projection lag below two seconds p95 at 25,000 synthetic
  events per minute.

## Measured recovery
The fictional backlog drained in 6 minutes, no events were lost, and duplicate projection
count remained zero.
"""
    (OUTPUT / ARTIFACTS[4][0]).write_text(incident, encoding="utf-8")

    adr = f"""# {NOTICE}

# Project Northstar ADR-004 — Revised Offline-Tolerant Event Flow

## Status
Accepted for the fictional hackathon demonstration only.

## Challenges that changed the draft
The design review identified replay ambiguity, missing row-level lineage, stale dashboard
visibility, dependency coupling, and an under-specified recovery gate. The synthetic
incident then demonstrated that total request latency could remain healthy while projection
freshness degraded.

## Revised decision
Use one validated command boundary, idempotency keys, a bounded durable projection queue,
versioned event envelopes, quarantine reason codes, source-batch lineage, optimistic
concurrency, dependency-specific readiness, and explicit freshness telemetry.

## Rejected alternatives
- Synchronous fan-out to every read model.
- Last-writer-wins updates for offline edits.
- Unbounded worker concurrency.
- A single unconditional health endpoint.

## Consequences and risks
The queue introduces observable lag and requires replay operations. Version conflicts need
an explicit user path. These costs are accepted because the components remain replaceable
and the entire exercise runs offline without Azure services.
"""
    (OUTPUT / ARTIFACTS[5][0]).write_text(adr, encoding="utf-8")

    lines = [
        f"# {NOTICE}",
        "",
        "# Project Northstar Measured Outcome and Retrospective",
        "",
        "All measurements below are invented outputs from deterministic synthetic tests.",
        "",
        "## Linked outcomes by role simulation",
    ]
    for record in records:
        decision = record["decisions"][3]
        lines.extend(
            [
                f"### {record['evidence_author']} — {record['evidence_role']}",
                f"{NOTICE} {decision['statement']}",
                f"- Decision: {decision['choice']}",
                f"- Measured synthetic outcome: {decision['outcome']}",
                f"- Remaining risk: {'; '.join(decision['risks'])}",
                "",
            ]
        )
    lines.extend(
        [
            "## Overall fictional result",
            "The final scripted run processed 25,000 events per minute at 1.6 seconds p95",
            "projection latency, reconciled every batch, and passed all 48 abuse cases.",
        ]
    )
    (OUTPUT / ARTIFACTS[6][0]).write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_audio_script(records: list[dict]) -> Path:
    path = OUTPUT / "04-architecture-review-audio-script.txt"
    lines = [
        NOTICE,
        "Project Northstar synthetic architecture review audio script.",
        "This recording is fictional and does not describe actual employee behavior.",
    ]
    for record in records:
        decision = record["decisions"][2]
        lines.append(
            f"{record['evidence_author']}, {record['evidence_role']}. "
            f"Synthetic decision: {decision['choice']}."
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _build_manifest(records: list[dict], docx_path: Path) -> None:
    docx_segments = parse_file(docx_path)
    comment_by_author = {}
    for segment in docx_segments:
        if segment.comment_id and segment.author not in comment_by_author:
            comment_by_author[segment.author] = segment
    chat_segments = {
        segment.author: segment
        for segment in parse_file(OUTPUT / ARTIFACTS[1][0])
        if segment.author
    }
    transcript_segments = {
        segment.author: segment
        for segment in parse_file(OUTPUT / ARTIFACTS[2][0])
        if segment.author
    }
    decisions = []
    for record in records:
        name = record["evidence_author"]
        role = record["evidence_role"]
        mappings = [
            (ARTIFACTS[0][0], record["decisions"][0], comment_by_author[name], "review_comment"),
            (ARTIFACTS[1][0], record["decisions"][1], chat_segments[name], "transcript_turn"),
            (
                ARTIFACTS[2][0],
                record["decisions"][2],
                transcript_segments[name],
                "transcript_turn",
            ),
        ]
        for artifact, decision, segment, evidence_type in mappings:
            decisions.append(
                {
                    "artifact": artifact,
                    "evidence_author": name,
                    "evidence_role": role,
                    "evidence_type": evidence_type,
                    "evidence_text": segment.text,
                    "page_or_segment": segment.page_or_segment,
                    "timestamp": segment.timestamp,
                    "comment_id": segment.comment_id,
                    "anchor_text": segment.anchor_text,
                    "parent_comment_id": segment.parent_comment_id,
                    "parent_author": segment.parent_author,
                    "parent_text": segment.parent_text,
                    **decision,
                }
            )
        outcome = record["decisions"][3]
        decisions.append(
            {
                "artifact": ARTIFACTS[6][0],
                "evidence_author": name,
                "evidence_role": role,
                "evidence_type": "document",
                "evidence_text": (
                    f"{NOTICE} {outcome['statement']} Decision: {outcome['choice']}. "
                    f"Measured synthetic outcome: {outcome['outcome']}"
                ),
                "page_or_segment": f"retrospective entry — {name}",
                "timestamp": None,
                "comment_id": None,
                "anchor_text": None,
                "parent_comment_id": None,
                "parent_author": None,
                "parent_text": None,
                **outcome,
            }
        )
    manifest = {
        "notice": NOTICE,
        "simulation_namespace": "northstar-role-sim-mixed-v2",
        "scenario": "Project Northstar",
        "artifacts": [
            {
                "filename": filename,
                "label": label,
                "source_type": source_type,
            }
            for filename, label, source_type in ARTIFACTS
        ],
        "decisions": decisions,
    }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def build() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    records = _profile_records()
    docx_path = _build_docx(records)
    _write_group_chat(records)
    _write_architecture_transcript(records)
    _write_markdown(records)
    _write_audio_script(records)
    _build_manifest(records, docx_path)


if __name__ == "__main__":
    build()
    print(f"Created mixed corpus in {OUTPUT}")
