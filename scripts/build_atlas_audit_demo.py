"""Build the synthetic AtlasAudit DOCX with real Word comments and reply threads."""

import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "demo-data" / "atlas-audit" / "01-initial-design-review.docx"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
W15_NS = "http://schemas.microsoft.com/office/word/2012/wordml"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
W = f"{{{W_NS}}}"
W14 = f"{{{W14_NS}}}"
W15 = f"{{{W15_NS}}}"

PARAGRAPHS = [
    (
        "Candidate platform",
        "The proposed hot path sends Event Hubs data to Azure Data Explorer for Kusto "
        "investigation. Synapse remains a candidate for long-range and cross-domain analysis.",
    ),
    (
        "Integrity proposal",
        "The first draft relies on normal storage access controls and daily archive export "
        "for audit retention.",
    ),
    (
        "Write topology",
        "A synchronous write to both Azure Data Explorer and Synapse could make both stores "
        "immediately available, but creates partial-success and replay coordination risk.",
    ),
    (
        "Operations",
        "The service will monitor ingestion failures and provide an operator dashboard.",
    ),
    (
        "Retention",
        "The initial estimate keeps 30 days of hot Azure Data Explorer data and 13 months "
        "in the archive.",
    ),
    (
        "Delivery",
        "The pilot has six weeks to prove 25,000 sustained and 80,000 burst events per "
        "second while meeting reliability, integrity, query, and cost gates.",
    ),
]

THREADS = [
    (
        1,
        2,
        1,
        "Daniel Kim",
        "Access control alone does not prove that exported evidence was not altered. Add "
        "source sequence ranges, immutable batches, and a hash-chain manifest.",
        "Maya Rao",
        "Accepted. Use hourly immutable Parquet batches with sequence ranges and chained "
        "manifest hashes, plus reconciliation for gaps and duplicates.",
    ),
    (
        3,
        4,
        2,
        "Priya Nair",
        "I suggest dual-writing to ADX and Synapse so either platform can answer immediately.",
        "Maya Rao",
        "Rejected for phase one. Synchronous dual-write increases partial-success and schema "
        "coordination risk. Keep one ADX ingress path and an idempotent archive export.",
    ),
    (
        5,
        6,
        3,
        "Luis Martinez",
        "The monitoring statement is too broad. Define searchable latency, backlog, oldest "
        "unexported batch, hash mismatch, dead-letter growth, and runbook ownership.",
        "Maya Rao",
        "Accepted. Add those signals and page when searchable latency exceeds 90 seconds or "
        "oldest-unexported-batch age exceeds ten minutes.",
    ),
    (
        7,
        8,
        4,
        "Aisha Khan",
        "Thirty hot days raises cost without evidence that incidents need that window. Model "
        "14 days and retain the 30-day estimate as a comparison baseline.",
        "Maya Rao",
        "Refined. Use 14-day hot retention, a 13-month immutable archive, and a pilot cost "
        "gate. Revisit only if measured incident demand shows the window is inadequate.",
    ),
    (
        9,
        10,
        5,
        "Ethan Brooks",
        "Please make the pilot gate testable: burst load, duplicate delivery, export restart, "
        "storage throttling, schema evolution, and archive recovery.",
        "Maya Rao",
        "Accepted. The decision remains phased until latency, failure recovery, integrity, "
        "and cost gates pass; a happy-path demonstration is not sufficient.",
    ),
]


def comment(comment_id: int, author: str, text: str, para_id: str) -> ET.Element:
    node = ET.Element(
        f"{W}comment",
        {
            f"{W}id": str(comment_id),
            f"{W}author": author,
            f"{W}date": f"2026-06-02T{9 + comment_id // 2:02d}:00:00Z",
        },
    )
    paragraph = ET.SubElement(node, f"{W}p", {f"{W14}paraId": para_id})
    run = ET.SubElement(paragraph, f"{W}r")
    ET.SubElement(run, f"{W}t").text = text
    return node


def add_reference(paragraph: ET.Element, comment_id: int) -> None:
    run = ET.Element(f"{W}r")
    ET.SubElement(run, f"{W}commentReference", {f"{W}id": str(comment_id)})
    paragraph.append(run)


def add_package_parts(parts: dict[str, bytes]) -> None:
    relationships_path = "word/_rels/document.xml.rels"
    relationships = ET.fromstring(parts[relationships_path])
    ET.SubElement(
        relationships,
        f"{{{REL_NS}}}Relationship",
        {
            "Id": "rIdAtlasComments",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/comments"
            ),
            "Target": "comments.xml",
        },
    )
    ET.SubElement(
        relationships,
        f"{{{REL_NS}}}Relationship",
        {
            "Id": "rIdAtlasCommentsExtended",
            "Type": (
                "http://schemas.microsoft.com/office/2011/"
                "relationships/commentsExtended"
            ),
            "Target": "commentsExtended.xml",
        },
    )
    parts[relationships_path] = ET.tostring(
        relationships, encoding="utf-8", xml_declaration=True
    )

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Override",
        {
            "PartName": "/word/comments.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.comments+xml"
            ),
        },
    )
    ET.SubElement(
        content_types,
        f"{{{CT_NS}}}Override",
        {
            "PartName": "/word/commentsExtended.xml",
            "ContentType": "application/vnd.ms-word.commentsExtended+xml",
        },
    )
    parts["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("AtlasAudit Initial Design Review", level=1)
    document.add_paragraph(
        "SYNTHETIC FICTIONAL HACKATHON DATA. AtlasAudit, Contoso Engineering, all "
        "people, metrics, comments, and decisions in this document are invented for "
        "demonstration and do not describe real events."
    )
    body_paragraphs = []
    for heading, text in PARAGRAPHS:
        document.add_heading(heading, level=2)
        body_paragraphs.append(document.add_paragraph(text))
    document.save(OUTPUT)

    with ZipFile(OUTPUT) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    document_xml = ET.fromstring(parts["word/document.xml"])
    xml_paragraphs = list(document_xml.iter(f"{W}p"))
    body_texts = [paragraph.text for paragraph in body_paragraphs]
    anchor_paragraphs = []
    for expected in body_texts:
        anchor_paragraphs.append(
            next(
                paragraph
                for paragraph in xml_paragraphs
                if "".join(node.text or "" for node in paragraph.iter(f"{W}t")) == expected
            )
        )

    comments = ET.Element(f"{W}comments")
    comments_extended = ET.Element(f"{W15}commentsEx")
    for (
        root_id,
        reply_id,
        anchor_index,
        root_author,
        root_text,
        reply_author,
        reply_text,
    ) in THREADS:
        paragraph = anchor_paragraphs[anchor_index]
        first_run = next(paragraph.iter(f"{W}r"))
        index = list(paragraph).index(first_run)
        paragraph.insert(index, ET.Element(f"{W}commentRangeStart", {f"{W}id": str(root_id)}))
        paragraph.insert(
            index + 2,
            ET.Element(f"{W}commentRangeEnd", {f"{W}id": str(root_id)}),
        )
        add_reference(paragraph, root_id)

        root_para_id = f"A{root_id:07X}"
        reply_para_id = f"B{reply_id:07X}"
        comments.append(comment(root_id, root_author, root_text, root_para_id))
        comments.append(comment(reply_id, reply_author, reply_text, reply_para_id))
        ET.SubElement(
            comments_extended,
            f"{W15}commentEx",
            {f"{W15}paraId": root_para_id},
        )
        ET.SubElement(
            comments_extended,
            f"{W15}commentEx",
            {
                f"{W15}paraId": reply_para_id,
                f"{W15}paraIdParent": root_para_id,
            },
        )

    parts["word/document.xml"] = ET.tostring(
        document_xml, encoding="utf-8", xml_declaration=True
    )
    parts["word/comments.xml"] = ET.tostring(
        comments, encoding="utf-8", xml_declaration=True
    )
    parts["word/commentsExtended.xml"] = ET.tostring(
        comments_extended, encoding="utf-8", xml_declaration=True
    )
    add_package_parts(parts)

    rebuilt = OUTPUT.with_suffix(".rebuilt.docx")
    with ZipFile(rebuilt, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    rebuilt.replace(OUTPUT)


if __name__ == "__main__":
    build()
    print(f"Created {OUTPUT}")
